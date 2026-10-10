"""Synthetic demo using real ingestion + BM25; never reads private corpus."""
import argparse,csv,json,os
from pathlib import Path
from docx import Document
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[1]
def main():
    parser=argparse.ArgumentParser();parser.add_argument('--ask',action='store_true',help='Optional paid LLM call; set API credentials first')
    args=parser.parse_args()
    config=json.loads((ROOT/'configs/demo.json').read_text(encoding='utf-8'))
    if not args.ask:
        os.environ.pop('DEEPSEEK_API_KEY',None)
        os.environ.pop('TYPESAFE_API_KEY',None)
    os.environ['RAG_RETRIEVAL_MODE']=config['retrieval_mode']
    os.environ['RAG_ENABLE_LLM_PLANNER']='0'
    os.environ['RAG_ORCHESTRATOR']='langgraph'
    os.environ['HAYSTACK_TELEMETRY_ENABLED']='false'
    from bank_rag.ingest import write_knowledge
    from bank_rag.pipeline import TrustedRAG
    from bank_rag.graph import build_rag_graph
    generated=ROOT/'data/demo/generated';raw=generated/'raw';raw.mkdir(parents=True,exist_ok=True)
    doc=Document()
    for line in (ROOT/'data/demo/mock_rules.md').read_text(encoding='utf-8').splitlines():
        if line.strip():doc.add_paragraph(line.lstrip('# '))
    doc.save(raw/'fictional_rules.docx')
    wb=Workbook();ws=wb.active;ws.title='虚构统计'
    with (ROOT/'data/demo/mock_statistics.csv').open(encoding='utf-8') as f:
        for row in csv.reader(f):ws.append(row)
    wb.save(raw/'fictional_statistics.xlsx')
    knowledge=ROOT/config['knowledge'];count=write_knowledge(raw,knowledge)
    # Make citations portable in generated synthetic records only.
    records=[json.loads(line) for line in knowledge.read_text(encoding='utf-8').splitlines()]
    for r in records:
        p=Path(r['meta'].get('source_path',''))
        if p.is_absolute():r['meta']['source_path']=p.relative_to(ROOT).as_posix()
    knowledge.write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in records),encoding='utf-8')
    rag=TrustedRAG(knowledge);graph=build_rag_graph(rag)
    results=[]
    for q in json.loads((ROOT/'data/demo/queries.json').read_text(encoding='utf-8')):
        docs=rag.retrieve(q);assert docs,'Demo retrieval returned no evidence'
        if args.ask:result=rag.ask(q)
        else:
            # Explicit demo-only presence check and extractive preview, not production Gate/LLM.
            result={'mode':'offline_evidence_preview','evidence_present':bool(docs),'answer_preview':docs[0].content,'citations':rag.citations(docs)}
        results.append({'question':q,'result':result})
    output={'indexed_records':count,'workflow_compiled':bool(graph),'llm_called':args.ask,'results':results}
    (generated/'demo_result.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(output,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
