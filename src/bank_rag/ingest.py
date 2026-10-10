from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable

from docx import Document as WordDocument
from openpyxl import load_workbook
import pandas as pd
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".xlsm"}


def _clean(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def _decode_legacy_text(data: bytes) -> str:
    """Decode text exported from legacy Chinese Word files without mojibake."""
    candidates: list[str] = []
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "utf-16"):
        try:
            candidates.append(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    if not candidates:
        return data.decode("gb18030", errors="replace")

    def quality(text: str) -> int:
        cjk = sum("\u4e00" <= char <= "\u9fff" for char in text)
        replacement = text.count("\ufffd")
        control = sum(ord(char) < 32 and char not in "\n\r\t" for char in text)
        return cjk * 3 - replacement * 100 - control * 20

    return max(candidates, key=quality)


def _chunk_legacy_text(content: str, chunk_size: int = 1200) -> list[str]:
    """Split a converted .doc at Chinese sentence boundaries for retrieval."""
    sentences = re.split(r"(?<=[\u3002\uff01\uff1f\uff1b])", content)
    chunks: list[str] = []
    buffer = ""
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        if buffer and len(buffer) + len(sentence) > chunk_size:
            chunks.append(buffer)
            buffer = ""
        # A rare unpunctuated line is still bounded for model context safety.
        while len(sentence) > chunk_size:
            if buffer:
                chunks.append(buffer)
                buffer = ""
            chunks.append(sentence[:chunk_size])
            sentence = sentence[chunk_size:]
        buffer += sentence
    if buffer:
        chunks.append(buffer)
    return chunks


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def resolve_source(source: Path, work_dir: Path) -> Path:
    """Return a readable source directory, extracting a ZIP once when needed."""
    if source.is_dir():
        return source
    if source.suffix.lower() != ".zip":
        raise ValueError("--source 必须是资料目录或 ZIP 数据包")
    extract_dir = work_dir / ".extracted"
    marker = extract_dir / ".source-name"
    if not extract_dir.exists() or not marker.exists() or marker.read_text(encoding="utf-8") != source.name:
        if extract_dir.exists():
            shutil.rmtree(extract_dir)
        extract_dir.mkdir(parents=True)
        with zipfile.ZipFile(source) as archive:
            archive.extractall(extract_dir)
        marker.write_text(source.name, encoding="utf-8")
    return extract_dir


def _base_meta(path: Path, source_root: Path) -> dict:
    relative = path.relative_to(source_root).as_posix()
    return {
        "doc_id": hashlib.sha1(relative.encode("utf-8")).hexdigest()[:16],
        "source_path": relative,
        "source_title": path.stem,
        "file_name": path.name,
        "file_type": path.suffix.lower().lstrip("."),
        "sha256": _sha256(path),
    }


def parse_pdf(path: Path, source_root: Path) -> list[dict]:
    meta = _base_meta(path, source_root)
    docs: list[dict] = []
    for page_number, page in enumerate(PdfReader(str(path)).pages, start=1):
        content = _clean(page.extract_text())
        if content:
            docs.append({"content": content, "meta": meta | {"page": page_number, "evidence_type": "page"}})
    return docs


def parse_docx(path: Path, source_root: Path) -> list[dict]:
    meta, docs, section, buffer = _base_meta(path, source_root), [], "正文", []
    document = WordDocument(str(path))
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = _clean(paragraph.text)
        if not text:
            continue
        if paragraph.style.name.lower().startswith("heading") or re.match(r"^(第[一二三四五六七八九十\d]+[章节条]|\d+(\.\d+)*)", text):
            if buffer:
                docs.append({"content": "\n".join(buffer), "meta": meta | {"section": section, "paragraph_end": index - 1, "evidence_type": "clause"}})
                buffer = []
            section = text
        buffer.append(text)
        if len("".join(buffer)) > 1000:
            docs.append({"content": "\n".join(buffer), "meta": meta | {"section": section, "paragraph_end": index, "evidence_type": "clause"}})
            buffer = []
    if buffer:
        docs.append({"content": "\n".join(buffer), "meta": meta | {"section": section, "evidence_type": "clause"}})
    docs.extend(_docx_table_records(path, document.tables, meta))
    return docs


def _docx_table_records(path: Path, tables: Iterable[object], meta: dict) -> list[dict]:
    """Append auditable Word table evidence without changing paragraph chunks.

    ``python-docx`` exposes the logical table grid through ``table.rows`` and
    ``row.cells``.  Each non-empty row and cell is retained, and each data row
    also receives a compact evidence group containing the leading header rows.
    This mirrors the existing spreadsheet evidence model while keeping DOCX
    paragraph ingestion unchanged.
    """
    docs: list[dict] = []
    for table_index, table in enumerate(tables, start=1):
        rows: list[list[str]] = []
        for row in table.rows:
            seen_merged_cells: set[int] = set()
            values: list[str] = []
            for cell in row.cells:
                cell_identity = id(cell._tc)
                if cell_identity in seen_merged_cells:
                    values.append("")
                    continue
                seen_merged_cells.add(cell_identity)
                values.append(_clean(cell.text))
            rows.append(values)
        non_empty_rows = [index for index, row in enumerate(rows) if any(row)]
        if not non_empty_rows:
            continue

        table_name = f"Word表格{table_index}"
        table_id = f"{meta['doc_id']}:docx:table:{table_index}"
        first_row = non_empty_rows[0]

        # Preserve up to three leading header rows.  The first row is always
        # retained as the table header; subsequent leading rows are included
        # only while they contain no standalone numeric value.
        header_rows = [first_row]
        for row_index in non_empty_rows[1:3]:
            if any(_is_numeric_cell(value) for value in rows[row_index]):
                break
            header_rows.append(row_index)

        def formatted_row(row_index: int) -> str:
            values = [
                f"{_excel_column(column_index + 1)}{row_index + 1}={value}"
                for column_index, value in enumerate(rows[row_index])
                if value
            ]
            return f"行 {row_index + 1}：" + "；".join(values)

        for row_index in non_empty_rows:
            row = rows[row_index]
            row_values = [
                f"{_excel_column(column_index + 1)}{row_index + 1}={value}"
                for column_index, value in enumerate(row)
                if value
            ]
            row_label = next((value for value in row if value), "")
            common_meta = meta | {
                "sheet": table_name,
                "table_index": table_index,
                "table_id": table_id,
                "row": row_index + 1,
                "row_label": row_label,
            }
            docs.append({
                "content": (
                    f"文件：{path.stem}\n表格：{table_index}\n行：{row_index + 1}\n"
                    + "；".join(row_values)
                ),
                "meta": common_meta | {"evidence_type": "table_row"},
            })

            for column_index, value in enumerate(row):
                if not value:
                    continue
                column = _excel_column(column_index + 1)
                cell = f"T{table_index}:R{row_index + 1}:C{column_index + 1}"
                header_label = " / ".join(
                    dict.fromkeys(
                        rows[header_index][column_index]
                        for header_index in header_rows
                        if column_index < len(rows[header_index]) and rows[header_index][column_index]
                    )
                )
                docs.append({
                    "content": (
                        f"文件：{path.stem}\n表格：{table_index}\n"
                        f"单元格：{cell}\n列：{column}\n"
                        f"表头：{header_label or '未识别'}\n数值：{value}"
                    ),
                    "meta": common_meta | {
                        "cell": cell,
                        "column": column_index + 1,
                        "column_label": header_label,
                        "evidence_type": "table_cell",
                    },
                })

            applicable_headers = [index for index in header_rows if index < row_index]
            content_parts = [f"文件：{path.stem}", f"表格：{table_index}"]
            if applicable_headers:
                content_parts.append("表头：\n" + "\n".join(formatted_row(index) for index in applicable_headers))
            content_parts.extend([
                "指标行：\n" + formatted_row(row_index),
                "数值单元格：\n" + "\n".join(
                    f"{_excel_column(column_index + 1)}{row_index + 1}｜表头："
                    + " / ".join(
                        dict.fromkeys(
                            rows[header_index][column_index]
                            for header_index in header_rows
                            if header_index < row_index
                            and column_index < len(rows[header_index])
                            and rows[header_index][column_index]
                        )
                    )
                    + f"｜数值：{value}"
                    for column_index, value in enumerate(row)
                    if value
                ),
            ])
            docs.append({
                "content": "\n".join(content_parts),
                "meta": common_meta | {
                    "header_rows": [index + 1 for index in applicable_headers],
                    "value_cells": [
                        f"{_excel_column(column_index + 1)}{row_index + 1}"
                        for column_index, value in enumerate(row)
                        if value
                    ],
                    "evidence_type": "table_evidence_group",
                },
            })
    return docs


def parse_doc(path: Path, source_root: Path) -> list[dict]:
    """Read legacy .doc with antiword or a headless LibreOffice conversion."""
    try:
        result = subprocess.run(["antiword", str(path)], capture_output=True, check=True)
        content = _clean(_decode_legacy_text(result.stdout))
    except (FileNotFoundError, subprocess.CalledProcessError):
        soffice_candidates = [
            os.getenv("SOFFICE_PATH"),
            shutil.which("soffice"),
            str(Path(os.getenv("ProgramFiles", "")) / "LibreOffice" / "program" / "soffice.exe"),
        ]
        soffice = next((candidate for candidate in soffice_candidates if candidate and Path(candidate).exists()), None)
        if soffice:
            output_dir = Path(tempfile.mkdtemp(prefix="bank_rag_doc_"))
            try:
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "txt:Text", "--outdir", str(output_dir), str(path)],
                    capture_output=True,
                    text=True,
                    timeout=120,
                    check=True,
                )
                text_path = output_dir / f"{path.stem}.txt"
                content = _clean(_decode_legacy_text(text_path.read_bytes())) if text_path.exists() else ""
            finally:
                # LibreOffice can keep a temporary lock briefly on Windows.
                # Cleanup must not turn a successfully extracted document into
                # a parse failure.
                shutil.rmtree(output_dir, ignore_errors=True)
            if content:
                meta = _base_meta(path, source_root)
                return [
                    {"content": chunk, "meta": meta | {"chunk": index, "evidence_type": "legacy_word_clause"}}
                    for index, chunk in enumerate(_chunk_legacy_text(content), start=1)
                ]
        if os.getenv("BANK_RAG_ENABLE_WORD_COM") != "1":
            raise RuntimeError(
                "未找到 antiword 或 LibreOffice。请安装其中之一，或仅在本机 Word 可无人值守启动时设置 BANK_RAG_ENABLE_WORD_COM=1。"
            )
        script = """
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
try {
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $document = $word.Documents.Open($env:BANK_RAG_DOC_PATH, $false, $true)
  [Console]::OutputEncoding = [System.Text.UTF8Encoding]::new()
  [Console]::Write($document.Content.Text)
} finally {
  if ($document) { $document.Close($false); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($document) }
  if ($word) { $word.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($word) }
}
"""
        environment = os.environ | {"BANK_RAG_DOC_PATH": str(path)}
        result = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=environment,
            timeout=90,
            check=True,
        )
        content = _clean(result.stdout)
    if not content:
        return []
    meta = _base_meta(path, source_root)
    return [
        {"content": chunk, "meta": meta | {"chunk": index, "evidence_type": "legacy_word_clause"}}
        for index, chunk in enumerate(_chunk_legacy_text(content), start=1)
    ]


def _nearest_label(values: list[list[object]], row: int, col: int) -> str:
    """Build a compact Excel label from the left and upper non-empty cells."""
    labels: list[str] = []
    for left in range(col - 1, max(-1, col - 4), -1):
        value = _clean(values[row][left])
        if value:
            labels.append(value)
            break
    for up in range(row - 1, max(-1, row - 6), -1):
        value = _clean(values[up][col])
        if value:
            labels.append(value)
            break
    return " | ".join(labels)


LARGE_SHEET_CELL_LIMIT = 50_000
LARGE_SHEET_ROW_GROUP_SIZE = 50


def _is_numeric_cell(value: object) -> bool:
    """Return whether one spreadsheet cell is a value suitable for evidence."""
    if isinstance(value, bool) or not _clean(value):
        return False
    if isinstance(value, (int, float)):
        return True
    text = _clean(value).replace(",", "")
    return bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?%?", text))


def _table_evidence_groups(
    path: Path,
    sheet_name: str,
    values: list[list[object]],
    meta: dict,
) -> Iterable[dict]:
    """Keep header, indicator row and value cells as one retrievable record.

    Row and cell records remain available for exact citations.  This grouped
    record prevents retrieval from separating a numeric cell from the header
    and indicator that define its meaning.
    """
    non_empty_rows: list[tuple[int, list[object]]] = [
        (row_index, row)
        for row_index, row in enumerate(values)
        if any(_clean(value) for value in row)
    ]
    for row_index, row in non_empty_rows:
        numeric_columns = [
            col_index for col_index, value in enumerate(row)
            if _is_numeric_cell(value)
        ]
        text_columns = [
            col_index for col_index, value in enumerate(row)
            if _clean(value) and not _is_numeric_cell(value)
        ]
        if not numeric_columns or not text_columns:
            continue

        # Retain the nearest preceding header-like rows. A header-like row has
        # text but no standalone numeric value; earlier data rows are excluded.
        header_rows: list[tuple[int, list[object]]] = []
        for candidate_index, candidate in reversed(non_empty_rows):
            if candidate_index >= row_index:
                continue
            if any(_is_numeric_cell(value) for value in candidate):
                continue
            header_rows.append((candidate_index, candidate))
            if len(header_rows) == 6:
                break
        header_rows.reverse()

        def formatted_row(index: int, source_row: list[object]) -> str:
            cells = [
                f"{_excel_column(col_index + 1)}{index + 1}={_clean(value)}"
                for col_index, value in enumerate(source_row)
                if _clean(value)
            ]
            return f"行 {index + 1}：" + "；".join(cells)

        value_lines: list[str] = []
        value_cells: list[str] = []
        for col_index in numeric_columns:
            cell = f"{_excel_column(col_index + 1)}{row_index + 1}"
            value_cells.append(cell)
            label = _nearest_label(values, row_index, col_index) or "未识别"
            value_lines.append(f"{cell}｜指标/口径：{label}｜数值：{_clean(row[col_index])}")

        indicator_label = _clean(row[text_columns[0]])
        content_parts = [f"文件：{path.stem}", f"工作表：{sheet_name}"]
        if header_rows:
            content_parts.append("表头：\n" + "\n".join(formatted_row(index, source_row) for index, source_row in header_rows))
        content_parts.extend([
            "指标行：\n" + formatted_row(row_index, row),
            "数值单元格：\n" + "\n".join(value_lines),
        ])
        yield {
            "content": "\n".join(content_parts),
            "meta": meta | {
                "sheet": sheet_name,
                "row": row_index + 1,
                "header_rows": [index + 1 for index, _row in header_rows],
                "value_cells": value_cells,
                "row_label": indicator_label,
                "evidence_type": "table_evidence_group",
            },
        }


def parse_xlsx(path: Path, source_root: Path) -> Iterable[dict]:
    """Keep cell citations for normal sheets; use row groups for very large templates."""
    meta = _base_meta(path, source_root)
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        # Avoid materializing every cell of template-sized sheets into Python memory.
        if sheet.max_row * sheet.max_column > LARGE_SHEET_CELL_LIMIT:
            group: list[tuple[int, list[str]]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                row_values = [f"{_excel_column(col_index + 1)}{row_index + 1}={_clean(value)}" for col_index, value in enumerate(row) if _clean(value)]
                if not row_values:
                    continue
                group.append((row_index, row_values))
                if len(group) == LARGE_SHEET_ROW_GROUP_SIZE:
                    content = "\n".join(f"行 {index + 1}：" + "；".join(values) for index, values in group)
                    yield {"content": f"文件：{path.stem}\n工作表：{sheet.title}\n{content}", "meta": meta | {"sheet": sheet.title, "row_start": group[0][0] + 1, "row_end": group[-1][0] + 1, "evidence_type": "table_row_group"}}
                    group = []
            if group:
                content = "\n".join(f"行 {index + 1}：" + "；".join(values) for index, values in group)
                yield {"content": f"文件：{path.stem}\n工作表：{sheet.title}\n{content}", "meta": meta | {"sheet": sheet.title, "row_start": group[0][0] + 1, "row_end": group[-1][0] + 1, "evidence_type": "table_row_group"}}
            continue
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        rows_with_content = [
            (row_index, [f"{_excel_column(col_index + 1)}{row_index + 1}={_clean(value)}" for col_index, value in enumerate(row) if _clean(value)])
            for row_index, row in enumerate(values)
        ]
        rows_with_content = [(row_index, row_values) for row_index, row_values in rows_with_content if row_values]
        if not rows_with_content:
            continue
        for row_index, row_values in rows_with_content:
            yield {"content": f"文件：{path.stem}\n工作表：{sheet.title}\n行：{row_index + 1}\n" + "；".join(row_values), "meta": meta | {"sheet": sheet.title, "row": row_index + 1, "evidence_type": "table_row"}}
        yield from _table_evidence_groups(path, sheet.title, values, meta)
        for row_index, row in enumerate(values):
            for col_index, value in enumerate(row):
                if not _clean(value):
                    continue
                cell = f"{_excel_column(col_index + 1)}{row_index + 1}"
                label = _nearest_label(values, row_index, col_index)
                content = f"文件：{path.stem}\n工作表：{sheet.title}\n单元格：{cell}\n指标/口径：{label or '未识别'}\n数值：{_clean(value)}"
                yield {"content": content, "meta": meta | {"sheet": sheet.title, "cell": cell, "row": row_index + 1, "column": col_index + 1, "evidence_type": "table_cell"}}


def _excel_column(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def enrich_table_metadata(record: dict) -> dict:
    """Store spreadsheet structure in metadata as well as in the evidence text."""
    meta = record["meta"]
    evidence_type = str(meta.get("evidence_type", ""))
    if not evidence_type.startswith("table_"):
        return record
    sheet = str(meta.get("sheet", ""))
    meta.setdefault("table_id", f"{meta['doc_id']}:{sheet}")
    cells = re.findall(r"([A-Z]+)\d+=([^\uFF1B;\n]+)", record["content"])
    labels = [value.strip() for _column, value in cells if value.strip() and not re.fullmatch(r"-?\d+(?:\.\d+)?", value.strip().replace(",", ""))]
    if labels and not meta.get("row_label"):
        meta["row_label"] = labels[0]
    return record


def parse_xls(path: Path, source_root: Path) -> list[dict]:
    """Parse official legacy .xls attachments via pandas/xlrd at row and cell level."""
    meta, docs = _base_meta(path, source_root), []
    workbook = pd.ExcelFile(path)
    for sheet_name in workbook.sheet_names:
        values = pd.read_excel(workbook, sheet_name=sheet_name, header=None).where(pd.notna, None).values.tolist()
        for row_index, row in enumerate(values):
            row_values = [f"{_excel_column(col_index + 1)}{row_index + 1}={_clean(value)}" for col_index, value in enumerate(row) if _clean(value)]
            if row_values:
                docs.append({"content": f"文件：{path.stem}\n工作表：{sheet_name}\n行：{row_index + 1}\n" + "；".join(row_values), "meta": meta | {"sheet": sheet_name, "row": row_index + 1, "evidence_type": "table_row"}})
        docs.extend(_table_evidence_groups(path, sheet_name, values, meta))
        for row_index, row in enumerate(values):
            for col_index, value in enumerate(row):
                if not _clean(value):
                    continue
                cell = f"{_excel_column(col_index + 1)}{row_index + 1}"
                label = _nearest_label(values, row_index, col_index)
                content = f"文件：{path.stem}\n工作表：{sheet_name}\n单元格：{cell}\n指标/口径：{label or '未识别'}\n数值：{_clean(value)}"
                docs.append({"content": content, "meta": meta | {"sheet": sheet_name, "cell": cell, "row": row_index + 1, "column": col_index + 1, "evidence_type": "table_cell"}})
    return docs


def parse_sources(source: Path, work_dir: Path) -> Iterable[dict]:
    source_root = resolve_source(source, work_dir)
    for path in source_root.rglob("*"):
        if path.name.startswith("._") or "__MACOSX" in path.parts or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        try:
            parser = {".pdf": parse_pdf, ".doc": parse_doc, ".docx": parse_docx, ".xls": parse_xls, ".xlsx": parse_xlsx, ".xlsm": parse_xlsx}[path.suffix.lower()]
            yield from parser(path, source_root)
        except Exception as exc:  # Preserve a trace instead of silently omitting a source.
            yield {"content": f"解析失败：{path.name}", "meta": _base_meta(path, source_root) | {"parse_error": str(exc), "evidence_type": "error"}}


def write_knowledge(source: Path, output: Path) -> int:
    output.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with output.open("w", encoding="utf-8") as handle:
        for record in parse_sources(source, output.parent):
            record = enrich_table_metadata(record)
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
            count += 1
    return count


def _is_weak_narrative(record: dict) -> bool:
    """Filter headings/contents pages that cannot independently support an answer."""
    content = _clean(record["content"])
    evidence_type = str(record["meta"].get("evidence_type", ""))
    if evidence_type == "error":
        return False
    if evidence_type.startswith("table_"):
        return False
    if len(content) < 80:
        return True
    compact = re.sub(r"[\s\d一二三四五六七八九十、.（）()：:-]", "", content)
    return len(compact) < 20


def _hierarchical_records(records: Iterable[dict]) -> Iterable[dict]:
    """Build a compact parent/child corpus for regulation and table retrieval.

    Narrative chunks keep their original locations but gain a stable parent id.
    A small parent summary supports document-level recall.  For spreadsheets we
    retain auditable row/group records and omit isolated cells, which otherwise
    dominate hybrid retrieval without helping deterministic table operations.
    """
    grouped: dict[str, list[dict]] = {}
    for record in records:
        grouped.setdefault(str(record["meta"]["doc_id"]), []).append(record)

    for doc_id, source_records in grouped.items():
        meta = source_records[0]["meta"]
        narrative = [
            record for record in source_records
            if not str(record["meta"].get("evidence_type", "")).startswith("table_")
            and record["meta"].get("evidence_type") != "error"
            and not _is_weak_narrative(record)
        ]
        if narrative:
            parent_id = f"{doc_id}:parent"
            sections = []
            for record in narrative:
                section = _clean(record["meta"].get("section", ""))
                if section and section not in sections:
                    sections.append(section)
            summary = f"文件：{meta['source_title']}\n文档类型：{meta['file_type']}"
            if sections:
                summary += "\n章节：" + "；".join(sections[:30])
            yield {
                "content": summary,
                "meta": meta | {
                    "parent_id": parent_id,
                    "hierarchy_level": "parent",
                    "child_count": len(narrative),
                    "evidence_type": "document_parent",
                },
            }
            for chunk_index, record in enumerate(narrative, start=1):
                yield {
                    "content": record["content"],
                    "meta": record["meta"] | {
                        "parent_id": parent_id,
                        "chunk_id": f"{doc_id}:chunk:{chunk_index}",
                        "chunk_index": chunk_index,
                        "hierarchy_level": "child",
                    },
                }

        for record in source_records:
            evidence_type = str(record["meta"].get("evidence_type", ""))
            if evidence_type == "error":
                yield record
            elif evidence_type in {"table_row", "table_row_group"}:
                yield {
                    "content": record["content"],
                    "meta": record["meta"] | {
                        "hierarchy_level": "table_row",
                        "table_schema": "row_cells_v1",
                    },
                }


def write_hierarchical_knowledge(source: Path, output: Path) -> int:
    """Write the V6 compact parent/child knowledge base without altering V5."""
    output.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with output.open("w", encoding="utf-8") as handle:
        for record in _hierarchical_records(parse_sources(source, output.parent)):
            record = enrich_table_metadata(record)
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
            count += 1
    return count
