from __future__ import annotations

import csv
import json
import math
import os
import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openai import OpenAI

from .pipeline import TrustedRAG


def _load_qa(qa_path: Path, limit: int | None = None, offset: int = 0) -> list[dict]:
    if qa_path.suffix.lower() == ".jsonl":
        records = [
            json.loads(line)
            for line in qa_path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        records = records[offset:]
        return records[:limit] if limit else records
    if qa_path.suffix.lower() == ".csv":
        with qa_path.open(encoding="utf-8-sig", newline="") as handle:
            records = list(csv.DictReader(handle))
        records = records[offset:]
        return records[:limit] if limit else records
    sheet = load_workbook(qa_path, read_only=True, data_only=True).active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    records = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    records = records[offset:]
    return records[:limit] if limit else records


def _question_with_options(item: dict) -> str:
    """Preserve the original multiple-choice choices when asking the RAG.

    The curated workbook keeps the stem in ``question`` and its A--D choices
    in separate columns.  Passing only the stem makes many regulatory records
    indistinguishable (for example, several rows all say "which statement is
    correct" while targeting different provisions in the same document).
    """
    question = str(item["question"])
    # Spreadsheet rows already use deterministic row/column operations. Their
    # numeric or textual distractors add no evidence and can be mistaken for
    # additional coordinates by a table parser, so preserve the original stem.
    if str(item.get("qa_type", "")).startswith("\u8868\u683c"):
        return question
    options = [
        (label, item.get(f"option_{label.lower()}"))
        for label in ("A", "B", "C", "D")
        if item.get(f"option_{label.lower()}") is not None
    ]
    if not options:
        return question
    rendered = "\n".join(f"{label}. {value}" for label, value in options)
    return f"{question}\n\n\u9009\u9879\uff1a\n{rendered}\n\n\u8bf7\u9009\u51fa\u6709\u8bc1\u636e\u652f\u6301\u7684\u6b63\u786e\u9009\u9879\uff0c\u5e76\u7ed9\u51fa\u9009\u9879\u5b57\u6bcd\u548c\u5b8c\u6574\u8868\u8ff0\u3002"


def _normalise(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value if value is not None else "")).lower()
    return re.sub(r"[\s\W_]+", "", text, flags=re.UNICODE)


def _numeric(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value)):
        return float(value)
    match = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*", str(value if value is not None else ""))
    return float(match.group(1)) if match else None


def answer_matches(expected: object, prediction: str) -> bool:
    """Deterministic scorer: tolerant number match, otherwise normalised containment."""
    expected_number = _numeric(expected)
    if expected_number is not None:
        candidates = re.findall(r"(?<![\w.])-?\d+(?:\.\d+)?", prediction.replace(",", ""))
        return any(math.isclose(float(candidate), expected_number, rel_tol=1e-5, abs_tol=1e-6) for candidate in candidates)
    expected_text = _normalise(expected)
    # Citations may be placed between two otherwise verbatim answer clauses;
    # they are provenance, not answer content, and must not break strict text
    # matching of a selected A--D option.
    prediction_without_citations = re.sub(r"\[E\d+\]", "", prediction, flags=re.IGNORECASE)
    return bool(expected_text) and expected_text in _normalise(prediction_without_citations)


def choice_matches(item: dict, prediction: str) -> bool | None:
    """Score explicit multiple-choice selection separately from answer wording."""
    expected = str(item.get("answer") or "").strip().upper()
    if expected not in {"A", "B", "C", "D"}:
        return None
    match = re.search(
        r"(?:\u6b63\u786e\u9009\u9879|\u552f\u4e00\u6b63\u786e\u9009\u9879|\u7b54\u6848|\u9009\u62e9)\s*(?:\u4e3a|\u662f)?\s*[:\uff1a]?\s*([A-D])",
        prediction,
        flags=re.IGNORECASE,
    )
    return match.group(1).upper() == expected if match else False


def _semantic_normalise(value: object) -> str:
    """Normalise harmless Chinese wording variants used in generated answers."""
    text = _normalise(value)
    for source, target in (("或者", "或"), ("遭到", "遭"), ("应当", "应"), ("属于", ""), ("该", "")):
        text = text.replace(source, target)
    return text


def _bigram_coverage(expected: str, actual: str) -> float:
    if len(expected) < 2:
        return float(expected in actual)
    bigrams = {expected[index : index + 2] for index in range(len(expected) - 1)}
    return sum(bigram in actual for bigram in bigrams) / len(bigrams)


def semantic_answer_matches(expected: object, prediction: str) -> bool:
    """Conservative text-semantic score; numeric answers remain exact."""
    if _numeric(expected) is not None:
        return answer_matches(expected, prediction)
    actual = _semantic_normalise(prediction)
    if not actual:
        return False
    # Multi-fact answers are separated by a semicolon in the supplied QA set.
    parts = [part.strip() for part in re.split(r"[;\uFF1B]", str(expected)) if _semantic_normalise(part)]
    for part in parts:
        required = _semantic_normalise(part)
        if required in actual:
            continue
        if _bigram_coverage(required, actual) < 0.86:
            return False
    return bool(parts)


def ai_judge_answer(expected: object, prediction: str) -> tuple[bool | None, str]:
    """Ask DeepSeek for a semantic equivalence judgement, separate from generation."""
    if not prediction.strip():
        return False, "empty_prediction"
    client = OpenAI(api_key=os.environ["DEEPSEEK_API_KEY"], base_url=os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com"))
    prompt = (
        "Judge whether the candidate answer is factually consistent with and answers the reference answer. "
        "Ignore wording, order, citations, and harmless rounding. Reject a different fact, incomplete multi-fact "
        "answer, refusal, or an answer that does not address the reference. Return exactly one line: "
        "VERDICT: YES or VERDICT: NO, followed by a short reason.\n\n"
        f"REFERENCE: {expected}\nCANDIDATE: {prediction}"
    )
    response = client.chat.completions.create(
        model=os.getenv("DEEPSEEK_MODEL", "deepseek-v4-flash"),
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
        max_tokens=120,
    )
    verdict = (response.choices[0].message.content or "").strip()
    return verdict.upper().startswith("VERDICT: YES"), verdict


def ai_review_answer(question: str, prediction: str, evidence: list[str]) -> tuple[str, str]:
    """Evidence-grounded review that deliberately does not see the gold answer."""
    if not prediction.strip():
        return "EMPTY", "empty_prediction"
    client = OpenAI(api_key=os.environ["DEEPSEEK_API_KEY"], base_url=os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com"))
    def build_prompt(max_chars_per_excerpt: int) -> str:
        excerpts = "\n\n".join(
            f"[E{index}] {text[:max_chars_per_excerpt]}" for index, text in enumerate(evidence, start=1)
        )
        return (
        "You are an evidence-grounded reviewer. Do NOT use or infer a gold answer. "
        "Judge only whether the candidate is a useful response to the question and is supported by the evidence. "
        "Return exactly one first line using one label: VERDICT: PASS (useful, supported factual answer), "
        "VERDICT: SAFE_REFUSAL (correctly explains evidence is insufficient), or VERDICT: FAIL (wrong, unsupported, "
        "irrelevant, or avoids an answer despite sufficient evidence). Then provide one short reason.\n\n"
        f"QUESTION: {question}\nCANDIDATE: {prediction}\nEVIDENCE:\n{excerpts}"
        )

    # Some DeepSeek responses can have an empty message body, especially with a
    # long evidence context. Retry once with a shorter context. An empty or
    # unparseable reviewer response is not evidence that the candidate failed.
    last_review = ""
    for max_chars_per_excerpt in (1200, 600):
        response = client.chat.completions.create(
            model=os.getenv("DEEPSEEK_MODEL", "deepseek-v4-flash"),
            messages=[{"role": "user", "content": build_prompt(max_chars_per_excerpt)}],
            temperature=0,
            max_tokens=120,
        )
        last_review = (response.choices[0].message.content or "").strip()
        match = re.search(r"VERDICT:\s*(PASS|SAFE_REFUSAL|FAIL)", last_review.upper())
        if match:
            return match.group(1), last_review
    return "UNAVAILABLE", last_review or "empty_review_after_retry"


def _source_hit(item: dict, citations: list[dict]) -> bool:
    expected_titles = item.get("source_titles")
    if expected_titles:
        expected = [_normalise(title) for title in expected_titles if _normalise(title)]
    else:
        expected = [_normalise(item.get("source_title"))]
    raw_label = str(item.get("file_label") or "")
    label_parts = [_normalise(part) for part in re.split(r"[;\uFF1B]", raw_label) if _normalise(part)]
    sources = " ".join(_normalise(citation.get("source_title")) for citation in citations)
    return bool(expected and all(title in sources for title in expected)) or bool(
        label_parts and all(label in sources for label in label_parts)
    )


def _rates(rows: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(rows)
    strict_correct = sum(row["strict_correct"] for row in rows)
    semantic_correct = sum(row["semantic_correct"] for row in rows)
    choice_rows = [row for row in rows if row.get("choice_correct") is not None]
    choice_correct = sum(bool(row["choice_correct"]) for row in choice_rows)
    judged_rows = [row for row in rows if row.get("ai_consistent") is not None]
    ai_judge_correct = sum(bool(row["ai_consistent"]) for row in judged_rows)
    review_attempted = [row for row in rows if row.get("ai_review_verdict")]
    reviewed_rows = [
        row for row in review_attempted
        if row["ai_review_verdict"] in {"PASS", "SAFE_REFUSAL", "FAIL"}
    ]
    review_unavailable = sum(row["ai_review_verdict"] == "UNAVAILABLE" for row in review_attempted)
    review_pass = sum(row["ai_review_verdict"] == "PASS" for row in reviewed_rows)
    safe_refusals = sum(row["ai_review_verdict"] == "SAFE_REFUSAL" for row in reviewed_rows)
    return {
        "questions": total,
        # Backward-compatible aliases: correct/accuracy mean strict score.
        "correct": strict_correct,
        "accuracy": round(strict_correct / total, 4) if total else 0,
        "strict_correct": strict_correct,
        "strict_accuracy": round(strict_correct / total, 4) if total else 0,
        "semantic_correct": semantic_correct,
        "semantic_accuracy": round(semantic_correct / total, 4) if total else 0,
        "choice_questions": len(choice_rows),
        "choice_correct": choice_correct,
        "choice_accuracy": round(choice_correct / len(choice_rows), 4) if choice_rows else None,
        "ai_judged": len(judged_rows),
        "ai_judge_correct": ai_judge_correct,
        "ai_judge_accuracy": round(ai_judge_correct / len(judged_rows), 4) if judged_rows else None,
        "ai_review_attempted": len(review_attempted),
        "ai_reviewed": len(reviewed_rows),
        "ai_review_unavailable": review_unavailable,
        "ai_review_pass": review_pass,
        "ai_review_useful_rate": round(review_pass / len(reviewed_rows), 4) if reviewed_rows else None,
        "ai_review_safe_refusals": safe_refusals,
        "ai_review_trusted_rate": round((review_pass + safe_refusals) / len(reviewed_rows), 4) if reviewed_rows else None,
        "source_hit_at_6": round(sum(row["source_hit"] for row in rows) / total, 4) if total else 0,
        "verified_rate": round(sum(row["verified"] for row in rows) / total, 4) if total else 0,
        "errors": sum(bool(row.get("error")) for row in rows),
    }


def _group_rates(rows: list[dict[str, Any]], field: str) -> dict[str, dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[row[field]].append(row)
    return {name: _rates(group) for name, group in groups.items()}


def evaluate_retrieval(qa_path: Path, rag: TrustedRAG, limit: int | None = None, offset: int = 0) -> dict:
    records = _load_qa(qa_path, limit, offset)
    rows = []
    for item in records:
        documents = rag.retrieve(str(item["question"]))
        citations = rag.citations(documents)
        rows.append({"correct": False, "strict_correct": False, "semantic_correct": False, "choice_correct": None, "ai_consistent": None, "source_hit": _source_hit(item, citations), "verified": False, "error": None, "qa_type": str(item["qa_type"]), "difficulty": str(item["difficulty_cn"])})
    return {"mode": "retrieval_only", **_rates(rows), "by_qa_type": _group_rates(rows, "qa_type"), "by_difficulty": _group_rates(rows, "difficulty")}


def evaluate_answers(qa_path: Path, rag: TrustedRAG, report_dir: Path, limit: int | None = None, offset: int = 0, judge: bool = False, review: bool = False) -> dict:
    """Ask every QA question and write an auditable per-question report plus summary."""
    if not os.getenv("DEEPSEEK_API_KEY"):
        raise RuntimeError("未配置 DEEPSEEK_API_KEY；请在运行评测的同一 PowerShell 窗口设置它。")
    report_dir.mkdir(parents=True, exist_ok=True)
    records = _load_qa(qa_path, limit, offset)
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(records, start=1):
        row: dict[str, Any] = {
            "id": item["id"],
            "qa_type": str(item["qa_type"]),
            "difficulty": str(item["difficulty_cn"]),
            "question": str(item["question"]),
            "options": {label: item.get(f"option_{label.lower()}") for label in ("A", "B", "C", "D")},
            "expected_answer": item["answer_text"],
            "prediction": "",
            "correct": False,
            "strict_correct": False,
            "semantic_correct": False,
            "choice_correct": None,
            "ai_consistent": None,
            "ai_judge_reason": "",
            "ai_review_verdict": "",
            "ai_review_reason": "",
            "answer_review": "",
            "generation_fallback": "",
            "source_hit": False,
            "verified": False,
            "citations": [],
            "error": None,
        }
        try:
            evaluation_question = _question_with_options(item)
            result = rag.ask(evaluation_question)
            row["prediction"] = result["answer"]
            row["citations"] = result["citations"]
            row["verified"] = bool(result["verified"])
            row["answer_review"] = str(result.get("answer_review") or "")
            row["generation_fallback"] = str(result.get("generation_fallback") or "")
            if result.get("generation_error"):
                row["error"] = str(result["generation_error"])
            row["strict_correct"] = answer_matches(row["expected_answer"], row["prediction"])
            row["semantic_correct"] = semantic_answer_matches(row["expected_answer"], row["prediction"])
            row["choice_correct"] = choice_matches(item, row["prediction"])
            row["correct"] = row["strict_correct"]
            row["source_hit"] = _source_hit(item, row["citations"])
            if judge:
                row["ai_consistent"], row["ai_judge_reason"] = ai_judge_answer(row["expected_answer"], row["prediction"])
            if review:
                documents = rag.retrieve(evaluation_question)
                row["ai_review_verdict"], row["ai_review_reason"] = ai_review_answer(
                    row["question"], row["prediction"], [document.content for document in documents]
                )
        except Exception as exc:  # A single API error must not discard the rest of the evaluation.
            row["error"] = f"{type(exc).__name__}: {exc}"
        rows.append(row)
        if index % 10 == 0 or index == len(records):
            print(f"已完成 {index}/{len(records)} 题")

    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "mode": "deepseek_answer_evaluation",
        "retrieval_mode": rag.retrieval_mode,
        **_rates(rows),
        "by_qa_type": _group_rates(rows, "qa_type"),
        "by_difficulty": _group_rates(rows, "difficulty"),
    }
    (report_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    with (report_dir / "details.jsonl").open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    with (report_dir / "details.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        fields = ["id", "qa_type", "difficulty", "question", "expected_answer", "prediction", "correct", "strict_correct", "semantic_correct", "choice_correct", "ai_consistent", "ai_judge_reason", "ai_review_verdict", "ai_review_reason", "answer_review", "generation_fallback", "source_hit", "verified", "error"]
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows({field: row[field] for field in fields} for row in rows)
    return summary
